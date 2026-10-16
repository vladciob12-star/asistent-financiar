"""Explicitly mapped CSV imports for external statements and reports.

No vendor schema is assumed here: mappings must name the exact headers in the
file being imported. This keeps Booking.com and bank imports auditable.
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from main import transforma_numar


@dataclass(frozen=True)
class FinancialRecord:
    source: str
    transaction_date: str
    description: str
    amount: str
    currency: str
    source_file: str
    source_row: int
    external_id: str = ""
    activity: str = "unassigned"
    kind: str = "unclassified"
    settlement: str = "recorded"
    counterparty: str = ""


def load_mapping(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Nu pot citi configurația de mapare {path}: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("Configurația de mapare trebuie să fie un obiect JSON.")
    required = {"source", "columns", "currency", "date_format"}
    missing = required - data.keys()
    if missing:
        raise ValueError("Lipsesc setările: " + ", ".join(sorted(missing)))
    columns = data["columns"]
    if not isinstance(columns, dict) or not {"date", "description", "amount"} <= columns.keys():
        raise ValueError("columns trebuie să indice exact antetele date, description și amount.")
    allowed_roles = {"date", "description", "amount", "external_id", "activity", "kind", "settlement", "counterparty"}
    unknown_roles = set(columns) - allowed_roles
    if unknown_roles:
        raise ValueError("Roluri de coloane necunoscute: " + ", ".join(sorted(unknown_roles)))
    for key in ("date", "description", "amount", "external_id", "activity", "kind", "settlement", "counterparty"):
        if key not in columns:
            continue
        if not isinstance(columns[key], str) or (key in {"date", "description", "amount"} and not columns[key].strip()):
            raise ValueError(f"columns.{key} trebuie să fie un antet nenul.")
    if "external_id" in columns and not isinstance(columns["external_id"], str):
        raise ValueError("columns.external_id trebuie să fie un antet CSV.")
    if not isinstance(data["source"], str) or not data["source"].strip():
        raise ValueError("source trebuie să fie un nume nenul, de exemplu banca sau Booking.com.")
    if not isinstance(data["currency"], str) or not re.fullmatch(r"[A-Za-z]{3}", data["currency"]):
        raise ValueError("currency trebuie să fie codul ISO 4217 de trei litere, furnizat explicit.")
    if not isinstance(data["date_format"], str) or not data["date_format"]:
        raise ValueError("date_format este obligatoriu și trebuie să corespundă formatului fișierului.")
    delimiter = data.get("delimiter", ",")
    if delimiter not in (",", ";", "\t", "|"):
        raise ValueError("delimiter acceptă doar virgulă, punct și virgulă, tab sau |.")
    data["delimiter"] = delimiter
    return data


def import_csv(path: Path, mapping: dict[str, Any]) -> list[FinancialRecord]:
    """Import records using only user-supplied header/date/currency mapping."""
    columns = mapping["columns"]
    try:
        stream = path.open("r", encoding="utf-8-sig", newline="")
    except OSError as exc:
        raise ValueError(f"Nu pot deschide {path}: {exc}") from exc
    records: list[FinancialRecord] = []
    with stream:
        reader = csv.DictReader(stream, delimiter=mapping["delimiter"])
        if not reader.fieldnames:
            raise ValueError(f"{path}: CSV-ul nu are antet.")
        reader.fieldnames = [name.strip().lstrip("\ufeff") for name in reader.fieldnames]
        header_map = {name: name for name in reader.fieldnames}
        for role, header in columns.items():
            if header not in header_map:
                raise ValueError(f"{path}: antetul mapat pentru {role!r} nu există: {header!r}")
        for line, row in enumerate(reader, start=2):
            if None in row:
                raise ValueError(f"{path}, rândul {line}: număr incorect de coloane.")
            try:
                date_text = (row.get(columns["date"]) or "").strip()
                transaction_date = datetime.strptime(date_text, mapping["date_format"]).date().isoformat()
            except ValueError as exc:
                raise ValueError(f"{path}, rândul {line}: data nu corespunde date_format.") from exc
            description = (row.get(columns["description"]) or "").strip()
            amount_text = (row.get(columns["amount"]) or "").strip()
            if not description or not amount_text:
                raise ValueError(f"{path}, rândul {line}: descrierea și suma sunt obligatorii.")
            try:
                amount = transforma_numar(amount_text, fisier=path, rand=line)
            except ValueError as exc:
                raise ValueError(str(exc)) from exc
            external_id = (row.get(columns.get("external_id", "")) or "").strip() if columns.get("external_id") else ""
            def mapped_text(role: str, default: str) -> str:
                header = columns.get(role)
                return (row.get(header) or "").strip() if header else default
            records.append(FinancialRecord(
                source=mapping["source"].strip(),
                transaction_date=transaction_date,
                description=description,
                amount=format(amount, "f"),
                currency=mapping["currency"].upper(),
                source_file=path.name,
                source_row=line,
                external_id=external_id,
                activity=mapped_text("activity", "unassigned"),
                kind=mapped_text("kind", "unclassified"),
                settlement=mapped_text("settlement", "recorded"),
                counterparty=mapped_text("counterparty", ""),
            ))
    if not records:
        raise ValueError(f"{path}: nu conține tranzacții.")
    return records


def import_xlsx(path: Path, mapping: dict[str, Any], *, sheet_name: str | None = None) -> list[FinancialRecord]:
    """Import a worksheet using the same explicit header mapping as CSV.

    Only the first row is treated as headers. Formulas are read as cached
    values; workbooks without cached values fail validation instead of guessing.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Pentru import Excel este necesar pachetul openpyxl.") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise ValueError(f"Nu pot deschide registrul Excel {path}: {exc}") from exc
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"{path}: foaia Excel nu există: {sheet_name!r}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        raw_headers = next(values, None)
        if not raw_headers:
            raise ValueError(f"{path}: foaia Excel nu are antet.")
        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        if len(set(headers)) != len(headers):
            raise ValueError(f"{path}: antete Excel duplicate.")
        columns = mapping["columns"]
        for role, header in columns.items():
            if header not in headers:
                raise ValueError(f"{path}: antetul mapat pentru {role!r} nu există: {header!r}")
        records: list[FinancialRecord] = []
        for line, cells in enumerate(values, start=2):
            row = dict(zip(headers, cells))
            if not any(value is not None and str(value).strip() for value in cells):
                continue
            try:
                date_value = row.get(columns["date"])
                if isinstance(date_value, datetime):
                    transaction_date = date_value.date().isoformat()
                elif isinstance(date_value, date):
                    transaction_date = date_value.isoformat()
                else:
                    transaction_date = datetime.strptime(str(date_value).strip(), mapping["date_format"]).date().isoformat()
            except (ValueError, TypeError, AttributeError) as exc:
                raise ValueError(f"{path}, rândul {line}: data nu corespunde date_format sau este goală.") from exc
            description = str(row.get(columns["description"]) or "").strip()
            amount_value = row.get(columns["amount"])
            if not description or amount_value is None:
                raise ValueError(f"{path}, rândul {line}: descrierea și suma sunt obligatorii.")
            try:
                amount = transforma_numar(str(amount_value), fisier=path, rand=line)
            except ValueError as exc:
                raise ValueError(str(exc)) from exc
            external_header = columns.get("external_id")
            external_value = row.get(external_header) if external_header else ""
            def mapped_text(role: str, default: str) -> str:
                header = columns.get(role)
                return str(row.get(header) or "").strip() if header else default
            records.append(FinancialRecord(
                source=mapping["source"].strip(), transaction_date=transaction_date,
                description=description, amount=format(amount, "f"),
                currency=mapping["currency"].upper(), source_file=path.name,
                source_row=line, external_id=str(external_value or "").strip(),
                activity=mapped_text("activity", "unassigned"),
                kind=mapped_text("kind", "unclassified"),
                settlement=mapped_text("settlement", "recorded"),
                counterparty=mapped_text("counterparty", "")))
        if not records:
            raise ValueError(f"{path}: nu conține tranzacții.")
        return records
    finally:
        workbook.close()


def export_records(records: list[FinancialRecord], path: Path) -> None:
    """Write normalized data and row-level provenance to a local CSV."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(asdict(records[0]).keys()))
        writer.writeheader()
        for record in records:
            row = asdict(record)
            for key in ("source", "description", "currency", "source_file", "external_id"):
                row[key] = _spreadsheet_safe(row[key])
            writer.writerow(row)


def _spreadsheet_safe(value: str) -> str:
    """Prevent imported text from being treated as a spreadsheet formula."""
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def monthly_totals(records: list[FinancialRecord]) -> list[dict[str, str]]:
    """Aggregate signed amounts by calendar month and currency, never across currencies."""
    totals: dict[tuple[str, str], Decimal] = {}
    for record in records:
        month = record.transaction_date[:7]
        key = (month, record.currency)
        totals[key] = totals.get(key, Decimal(0)) + Decimal(record.amount)
    return [{"luna": month, "moneda": currency, "total_semnat": format(amount, ".2f")}
            for (month, currency), amount in sorted(totals.items())]


def export_monthly_totals(records: list[FinancialRecord], path: Path) -> None:
    rows = monthly_totals(records)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["luna", "moneda", "total_semnat"])
        writer.writeheader()
        writer.writerows(rows)
