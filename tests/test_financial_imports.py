import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from financial_imports import FinancialRecord, export_records, import_csv, import_mt940, import_xlsx, load_mapping, monthly_totals
from main import citeste_raport, transforma_numar


class NumberParsingTests(unittest.TestCase):
    def test_romanian_and_plain_decimal_formats(self):
        self.assertEqual(transforma_numar("1.234,50 lei"), 1234.50)
        self.assertEqual(transforma_numar("1234.50"), 1234.50)

    def test_invalid_number_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Număr invalid"):
            transforma_numar("12x")

    def test_non_finite_number_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Număr invalid"):
            transforma_numar("NaN")


class CreasoftTests(unittest.TestCase):
    def test_aggregates_by_vendomat(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sales.csv"
            path.write_text(
                "Vendomat;Cantitate;Incasari totale (cash + non-cash - rest)\n"
                "A;2;10,25\nA;3;4,75\nB;1;2.00\n", encoding="utf-8")
            totals = citeste_raport(path)
        self.assertEqual(totals["A"]["cantitate"], 5)
        self.assertEqual(totals["A"]["incasari"], 15)

    def test_missing_creasoft_column_is_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "bad.csv"
            path.write_text("Vendomat;Cantitate\nA;1\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "lipsesc coloanele"):
                citeste_raport(path)


class MappedCsvTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.mapping = {
            "source": "extras-importate",
            "currency": "RON",
            "date_format": "%d.%m.%Y",
            "delimiter": ";",
            "columns": {"date": "Valuta date", "description": "Detalii", "amount": "Suma semnata", "external_id": "ID"},
        }
        self.csv_path = self.folder / "statement.csv"
        self.csv_path.write_text("Valuta date;Detalii;Suma semnata;ID\n01.10.2026;Plata furnizor;-25,50;T-1\n",
                                 encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def test_import_uses_explicit_mapping_and_keeps_provenance(self):
        record = import_csv(self.csv_path, self.mapping)[0]
        self.assertEqual(record, FinancialRecord("extras-importate", "2026-10-01", "Plata furnizor",
                                                  "-25.50", "RON", "statement.csv", 2, "T-1"))

    def test_missing_real_header_fails_closed(self):
        self.mapping["columns"]["amount"] = "Invented Amount Header"
        with self.assertRaisesRegex(ValueError, "nu există"):
            import_csv(self.csv_path, self.mapping)

    def test_invalid_date_is_rejected_with_row_number(self):
        self.csv_path.write_text("Valuta date;Detalii;Suma semnata;ID\nnot-a-date;Plata;-2;T-1\n", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "rândul 2"):
            import_csv(self.csv_path, self.mapping)

    def test_mapping_file_requires_currency_and_column_headers(self):
        path = self.folder / "mapping.json"
        path.write_text(json.dumps(self.mapping), encoding="utf-8")
        self.assertEqual(load_mapping(path)["currency"], "RON")
        self.mapping["currency"] = ""
        path.write_text(json.dumps(self.mapping), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "ISO 4217"):
            load_mapping(path)

    def test_export_contains_normalized_rows(self):
        records = import_csv(self.csv_path, self.mapping)
        output = self.folder / "out" / "normalized.csv"
        export_records(records, output)
        with output.open(encoding="utf-8-sig", newline="") as stream:
            saved = list(csv.DictReader(stream))
        self.assertEqual(saved[0]["external_id"], "T-1")
        self.assertEqual(saved[0]["source_row"], "2")

    def test_monthly_totals_keep_currencies_separate(self):
        records = [
            FinancialRecord("bank", "2026-10-01", "in", "100", "RON", "a.csv", 2),
            FinancialRecord("bank", "2026-10-02", "out", "-20", "RON", "a.csv", 3),
            FinancialRecord("booking", "2026-10-02", "payout", "50", "EUR", "b.csv", 2),
        ]
        self.assertEqual(monthly_totals(records), [
            {"luna": "2026-10", "moneda": "EUR", "total_semnat": "50.00"},
            {"luna": "2026-10", "moneda": "RON", "total_semnat": "80.00"},
        ])

    def test_export_neutralizes_spreadsheet_formulas_in_text(self):
        record = FinancialRecord("bank", "2026-10-01", "=HYPERLINK(\"https://bad\")", "2.00", "RON", "x.csv", 2)
        output = self.folder / "safe.csv"
        export_records([record], output)
        with output.open(encoding="utf-8-sig", newline="") as stream:
            saved = next(csv.DictReader(stream))
        self.assertTrue(saved["description"].startswith("'="))

    def test_excel_import_uses_explicit_mapping_and_accepts_date_cells(self):
        from openpyxl import Workbook
        path=self.folder/"statement.xlsx"
        workbook=Workbook()
        sheet=workbook.active
        sheet.title="Extrase"
        sheet.append(["Valuta date","Detalii","Suma semnata","ID"])
        from datetime import datetime
        sheet.append([datetime(2026,10,1),"Comision",-1.25,"F-1"])
        workbook.save(path)
        workbook.close()
        rows=import_xlsx(path,self.mapping,sheet_name="Extrase")
        self.assertEqual(rows[0].transaction_date,"2026-10-01")
        self.assertEqual(rows[0].amount,"-1.25")
        self.assertEqual(rows[0].external_id,"F-1")

    def test_excel_import_rejects_missing_mapped_headers(self):
        from openpyxl import Workbook
        path=self.folder/"bad.xlsx"
        workbook=Workbook()
        workbook.active.append(["Date","Description","Amount"])
        workbook.active.append(["01.10.2026","Test",1])
        workbook.save(path)
        workbook.close()
        with self.assertRaisesRegex(ValueError,"antetul mapat"):
            import_xlsx(path,self.mapping)

    def test_mt940_import_parses_synthetic_transactions_and_hashes_reference(self):
        path = self.folder / "synthetic.sta"
        path.write_text(":20:SYNTHETIC\n:25:ACCOUNT-REDACTED\n:60F:C261001RON1000,00\n"
                        ":61:2610021002C100,50NTRFNONREF//SYNTH-REF-1\n:86:INCASARE TEST\n"
                        ":61:261003D25,50NTRFNONREF//SYNTH-REF-2\n:86:PLATA TEST\n"
                        ":62F:C261003RON1075,00\n", encoding="ascii")
        rows = import_mt940(path)
        self.assertEqual([(row.transaction_date, row.amount, row.currency) for row in rows], [
            ("2026-10-02", "100.50", "RON"), ("2026-10-03", "-25.50", "RON")])
        self.assertEqual(rows[0].description, "INCASARE TEST")
        self.assertEqual(len(rows[0].external_id), 64)
        self.assertNotIn("SYNTH-REF-1", rows[0].external_id)

    def test_mt940_rejects_statement_without_currency(self):
        path = self.folder / "missing-currency.sta"
        path.write_text(":20:SYNTHETIC\n:61:261002C100,00NTRFNONREF\n", encoding="ascii")
        with self.assertRaisesRegex(ValueError, "nu declară moneda"):
            import_mt940(path)

    def test_cli_excel_import_exports_normalized_ledger(self):
        from openpyxl import Workbook
        path=self.folder/"input.xlsx"
        workbook=Workbook()
        sheet=workbook.active
        sheet.append(["Valuta date","Detalii","Suma semnata","ID"])
        from datetime import datetime
        sheet.append([datetime(2026,10,1),"Test expense",-2.5,"X-2"])
        workbook.save(path)
        workbook.close()
        mapping_path=self.folder/"mapping.json"
        mapping_path.write_text(json.dumps(self.mapping),encoding="utf-8")
        output=self.folder/"out.csv"
        monthly=self.folder/"months.csv"
        main=Path(__file__).resolve().parents[1]/"main.py"
        result=subprocess.run([sys.executable,str(main),"--import-file",str(path),"--mapping",str(mapping_path),
            "--export",str(output),"--monthly-export",str(monthly)],capture_output=True,text=True,check=False)
        self.assertEqual(result.returncode,0,result.stderr)
        with output.open(encoding="utf-8-sig",newline="") as stream:
            saved=list(csv.DictReader(stream))
        self.assertEqual(saved[0]["external_id"],"X-2")
        self.assertEqual(saved[0]["amount"],"-2.5")
        self.assertTrue(monthly.exists())

    def test_cli_booking_csv_stores_scheduled_payout_as_receivable(self):
        from finance_store import connect, monthly_summary
        path=self.folder/"booking.csv"
        path.write_text("Data;Detalii;Suma;Rezervare;Tip;Stare;Activitate\n"
                        "01.10.2026;Payout programat;400.00;R-1;income;scheduled;booking_property\n",encoding="utf-8")
        mapping={"source":"Booking.com","currency":"RON","date_format":"%d.%m.%Y","delimiter":";",
            "columns":{"date":"Data","description":"Detalii","amount":"Suma","external_id":"Rezervare",
                       "kind":"Tip","settlement":"Stare","activity":"Activitate"}}
        mapping_path=self.folder/"booking-map.json"
        mapping_path.write_text(json.dumps(mapping),encoding="utf-8")
        database=self.folder/"ledger.sqlite3"
        result=subprocess.run([sys.executable,str(Path(__file__).resolve().parents[1]/"main.py"),"--import-file",str(path),
            "--mapping",str(mapping_path),"--export",str(self.folder/"booking-out.csv"),"--database",str(database)],
            capture_output=True,text=True,check=False)
        self.assertEqual(result.returncode,0,result.stderr)
        db=connect(database)
        try:
            summary=monthly_summary(db,"2026-10")
        finally:
            db.close()
        self.assertEqual(summary["booking_receivables_not_cash"],[{"currency":"RON","amount":"400.00"}])
        self.assertEqual(summary["activities"],{})


if __name__ == "__main__":
    unittest.main()
